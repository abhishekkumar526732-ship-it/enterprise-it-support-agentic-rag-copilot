from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}



def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()
    if suffix in {".txt", ".md"}:
        return TextLoader(str(path), encoding="utf-8").load()
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")



def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))


def upload_to_s3_if_enabled(local_path: Path, object_name: str | None = None) -> bool:
    """
    If AWS S3 is enabled, mirrors the ingested file to an S3 bucket
    for durable, centralized storage across stateless EC2 instances.
    """
    from app.core.config import get_settings
    settings = get_settings()
    if not settings.s3_enabled or not settings.s3_bucket_name:
        return False

    key = object_name or local_path.name
    try:
        import boto3
        s3 = boto3.client("s3", region_name=settings.aws_region)
        s3.upload_file(str(local_path), settings.s3_bucket_name, key)
        return True
    except Exception:
        return False
